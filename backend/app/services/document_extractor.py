import csv
import io
import os
import uuid
from pathlib import Path
from typing import Dict, Any, Tuple
from werkzeug.utils import secure_filename
from backend.app.utils.errors import ValidationError, BadRequestError


class DocumentExtractor:
    @staticmethod
    def validate_file(filename: str, file_size: int, allowed_extensions: set, max_size_bytes: int):
        if not filename or "." not in filename:
            raise ValidationError("File must have a valid extension.")
        ext = filename.rsplit(".", 1)[1].lower()
        if ext not in allowed_extensions:
            raise ValidationError(
                f"File extension '.{ext}' is not supported. Allowed extensions: {', '.join(sorted(allowed_extensions))}"
            )
        if file_size > max_size_bytes:
            raise ValidationError(
                f"File size exceeds limit of {max_size_bytes // (1024 * 1024)} MB."
            )
        return ext

    @staticmethod
    def save_file_safely(file_storage, upload_dir: str) -> Tuple[str, str, int]:
        """
        Saves uploaded file securely into upload_dir, preventing directory traversal.
        Returns: (original_filename, stored_filename, file_size)
        """
        original_name = secure_filename(file_storage.filename or "uploaded_file")
        ext = original_name.rsplit(".", 1)[-1].lower() if "." in original_name else "txt"
        unique_id = str(uuid.uuid4())[:8]
        stored_name = f"{unique_id}_{original_name}"

        dest_dir = Path(upload_dir).resolve()
        dest_path = (dest_dir / stored_name).resolve()

        # Prevent path traversal
        if not str(dest_path).startswith(str(dest_dir)):
            raise BadRequestError("Invalid file path detected.")

        dest_dir.mkdir(parents=True, exist_ok=True)
        file_storage.save(str(dest_path))
        file_size = os.path.getsize(dest_path)

        return original_name, stored_name, file_size

    @classmethod
    def extract_content(cls, file_path: str, ext: str) -> Tuple[str, Dict[str, Any]]:
        """Extracts text and structured metadata based on file extension."""
        ext = ext.lower()
        if ext == "txt":
            return cls._extract_txt(file_path)
        elif ext == "csv":
            return cls._extract_csv(file_path)
        elif ext in ("xlsx", "xls"):
            return cls._extract_xlsx(file_path)
        elif ext == "pdf":
            return cls._extract_pdf(file_path)
        elif ext == "docx":
            return cls._extract_docx(file_path)
        else:
            return cls._extract_txt(file_path)

    @staticmethod
    def _extract_txt(file_path: str) -> Tuple[str, Dict[str, Any]]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        lines = text.splitlines()
        metadata = {
            "type": "txt",
            "line_count": len(lines),
            "word_count": len(text.split()),
            "character_count": len(text),
        }
        return text, metadata

    @staticmethod
    def _extract_csv(file_path: str) -> Tuple[str, Dict[str, Any]]:
        rows = []
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append(row)

        header = rows[0] if rows else []
        data_rows = rows[1:] if len(rows) > 1 else []
        preview_text = f"CSV with {len(data_rows)} rows and {len(header)} columns: {', '.join(header[:10])}\n\n"
        for r in data_rows[:15]:
            preview_text += " | ".join(str(cell) for cell in r) + "\n"

        metadata = {
            "type": "csv",
            "columns": header,
            "row_count": len(data_rows),
            "sample_rows": data_rows[:5],
        }
        return preview_text, metadata

    @staticmethod
    def _extract_xlsx(file_path: str) -> Tuple[str, Dict[str, Any]]:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheets_data = {}
            full_text_lines = []

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                header = [str(c or "") for c in rows[0]] if rows else []
                sheet_rows = len(rows) - 1 if len(rows) > 1 else 0
                sheets_data[sheet_name] = {
                    "columns": header[:20],
                    "row_count": sheet_rows,
                }
                full_text_lines.append(f"=== Sheet: {sheet_name} ({sheet_rows} rows) ===")
                for r in rows[:10]:
                    full_text_lines.append(" | ".join(str(c or "") for c in r[:15]))
                full_text_lines.append("\n")

            full_text = "\n".join(full_text_lines)
            metadata = {
                "type": "xlsx",
                "sheets": wb.sheetnames,
                "sheet_details": sheets_data,
            }
            return full_text, metadata
        except Exception as exc:
            return f"Error reading XLSX file: {str(exc)}", {"type": "xlsx", "error": str(exc)}

    @staticmethod
    def _extract_pdf(file_path: str) -> Tuple[str, Dict[str, Any]]:
        try:
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            num_pages = len(reader.pages)
            text_parts = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                text_parts.append(f"--- Page {i + 1} ---\n{page_text}")

            full_text = "\n\n".join(text_parts)
            metadata = {
                "type": "pdf",
                "page_count": num_pages,
                "word_count": len(full_text.split()),
            }
            return full_text, metadata
        except Exception as exc:
            return f"Error reading PDF file: {str(exc)}", {"type": "pdf", "error": str(exc)}

    @staticmethod
    def _extract_docx(file_path: str) -> Tuple[str, Dict[str, Any]]:
        try:
            import docx

            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            metadata = {
                "type": "docx",
                "paragraph_count": len(paragraphs),
                "word_count": len(full_text.split()),
            }
            return full_text, metadata
        except Exception as exc:
            return f"Error reading DOCX file: {str(exc)}", {"type": "docx", "error": str(exc)}

